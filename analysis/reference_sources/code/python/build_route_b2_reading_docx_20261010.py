"""Format the existing reading manuscript as Word; do not revise its science."""
import hashlib
import json
import re
from pathlib import Path
from urllib.parse import unquote
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'manuscript/Route_B2_scientific_reading_draft_en.md'
OUTPUT = ROOT / 'manuscript/Route_B2_scientific_reading_draft_en_20261010_v3.docx'
QA = ROOT / 'results/diagnostics/reading_docx_20261010_v3'
SOURCE_SHA = '13ad6027ead6bba808fc6e3481474d86fe73d43cf71844c13da30fd80e2d79bd'
INLINE = re.compile(r'\[([^\]\n]+)\]\(([^)\n]+)\)|\*\*([^*]+)\*\*|`([^`]+)`|\*([^*\n]+)\*')
MATH = re.compile(r'Cov\(D,T\) = −v_SNI|Corr\(D,T\) = −v_SNI / sqrt\[\(v_SNI \+ v_sham\)\(v_PLM \+ v_SNI\)\]')
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'


def plain(text):
    return INLINE.sub(lambda m: next(v for v in (m[1], m[3], m[4], m[5]) if v is not None), text)


def math_run(text):
    r = OxmlElement('m:r')
    t = OxmlElement('m:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    return r


def variance(label):
    s = OxmlElement('m:sSub')
    for name, child in [('e', math_run('v')), ('sub', math_run(label))]:
        el = OxmlElement('m:' + name)
        el.append(child)
        s.append(el)
    return s


def equation(text):
    out = OxmlElement('m:oMath')
    if text.startswith('Cov'):
        out.extend([math_run('Cov(D,T) = −'), variance('SNI')])
    else:
        out.append(math_run('Corr(D,T) = '))
        fraction = OxmlElement('m:f')
        num, den = OxmlElement('m:num'), OxmlElement('m:den')
        num.extend([math_run('−'), variance('SNI')])
        rad = OxmlElement('m:rad')
        prop = OxmlElement('m:radPr')
        hidden = OxmlElement('m:degHide')
        hidden.set(qn('m:val'), '1')
        prop.append(hidden)
        rad.extend([prop, OxmlElement('m:deg')])
        e = OxmlElement('m:e')
        for a, b in [('SNI', 'sham'), ('PLM', 'SNI')]:
            delim = OxmlElement('m:d')
            dp = OxmlElement('m:dPr')
            for name, value in [('begChr', '('), ('endChr', ')')]:
                tag = OxmlElement('m:' + name)
                tag.set(qn('m:val'), value)
                dp.append(tag)
            de = OxmlElement('m:e')
            de.extend([variance(a), math_run(' + '), variance(b)])
            delim.extend([dp, de])
            e.append(delim)
        rad.append(e)
        den.append(rad)
        fraction.extend([num, den])
        out.append(fraction)
    return out


def visible_xml(el):
    tag = el.tag
    if tag in (W + 't', M + 't'):
        return el.text or ''
    if tag == W + 'br':
        return '\n'
    if tag == M + 'sSub':
        return visible_xml(el.find(M + 'e')) + '_' + visible_xml(el.find(M + 'sub'))
    if tag == M + 'f':
        return visible_xml(el.find(M + 'num')) + ' / ' + visible_xml(el.find(M + 'den'))
    if tag == M + 'rad':
        return 'sqrt[' + visible_xml(el.find(M + 'e')) + ']'
    if tag == M + 'd':
        return '(' + visible_xml(el.find(M + 'e')) + ')'
    if tag.endswith('Pr'):
        return ''
    return ''.join(visible_xml(c) for c in el)


links = []


def add_inline(paragraph, text):
    cursor = 0
    matches = sorted(list(INLINE.finditer(text)) + list(MATH.finditer(text)), key=lambda m: m.start())
    for m in matches:
        if m.start() < cursor:
            continue
        paragraph.add_run(text[cursor:m.start()])
        if m.re is MATH:
            paragraph._p.append(equation(m[0]))
        elif m[1] is not None:
            label, target = m[1], m[2]
            if not target.startswith(('https://', 'http://', 'mailto:')):
                path, _, fragment = target.partition('#')
                full = (SOURCE.parent / unquote(path)).resolve()
                assert full.exists(), target
                target = full.as_uri() + ('#' + fragment if fragment else '')
            rel = paragraph.part.relate_to(target, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
            hyperlink = OxmlElement('w:hyperlink')
            hyperlink.set(qn('r:id'), rel)
            r = OxmlElement('w:r')
            prop = OxmlElement('w:rPr')
            color = OxmlElement('w:color')
            color.set(qn('w:val'), '1F497D')
            underline = OxmlElement('w:u')
            underline.set(qn('w:val'), 'single')
            prop.extend([color, underline])
            t = OxmlElement('w:t')
            t.text = label
            r.extend([prop, t])
            hyperlink.append(r)
            paragraph._p.append(hyperlink)
            links.append({'label': label, 'target': target})
        else:
            run = paragraph.add_run(next(v for v in (m[3], m[4], m[5]) if v is not None))
            run.bold = m[3] is not None
            run.italic = m[5] is not None
            if m[4] is not None:
                run.font.name = 'Consolas'
                run.font.size = Pt(10)
        cursor = m.end()
    paragraph.add_run(text[cursor:])


def build():
    source = SOURCE.read_bytes()
    assert hashlib.sha256(source).hexdigest() == SOURCE_SHA
    assert not OUTPUT.exists() and not QA.exists(), 'Preserve previous outputs'
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.72)
    section.left_margin = section.right_margin = Inches(0.72)
    for style in doc.styles:
        if style.type == 1:
            style.font.name = 'Times New Roman'
            style.font.size = Pt(11.5)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'SimSun')
    normal = doc.styles['Normal'].paragraph_format
    normal.line_spacing = 1.12
    normal.space_after = Pt(6)
    normal.widow_control = True
    for name, size in [('Title', 18), ('Heading 1', 14), ('Heading 2', 12), ('Heading 3', 11.5)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = name != 'Title'
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    r, t = OxmlElement('w:r'), OxmlElement('w:t')
    t.text = '1'
    r.append(t)
    field.append(r)
    footer._p.append(field)
    expected, heading_map, table_n, reference_n = [], [], 0, 0
    lines = source.decode('utf-8').splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line.startswith('|'):
            rows = [line]
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i].strip())
                i += 1
            cells = [[x.strip() for x in row.strip('|').split('|')] for row in rows]
            assert len(cells) >= 3 and all(re.fullmatch(r':?-+:?', c) for c in cells[1])
            cells.pop(1)
            table_n += 1
            widths = {3: [1.2, 2.55, 3.08], 6: [0.9, 1.14, 1.14, 1.12, 1.36, 1.17], 4: [0.85, 1.0, 2.38, 2.6]}[len(cells[0])]
            table = doc.add_table(rows=0, cols=len(cells[0]))
            table.autofit = False
            for column, width in zip(table.columns, widths):
                column.width = Inches(width)
            borders = OxmlElement('w:tblBorders')
            for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                border = OxmlElement('w:' + side)
                for k, v in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    border.set(qn('w:' + k), v)
                borders.append(border)
            table._tbl.tblPr.append(borders)
            visible_rows = []
            for row_i, row in enumerate(cells):
                assert len(row) == len(widths)
                created = table.add_row()
                created._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                if row_i == 0:
                    repeat = OxmlElement('w:tblHeader')
                    created._tr.get_or_add_trPr().append(repeat)
                for col_i, (cell, value, width) in enumerate(zip(created.cells, row, widths)):
                    cell.width = Inches(width)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    padding = OxmlElement('w:tcMar')
                    for side in ('top', 'bottom', 'left', 'right'):
                        side_el = OxmlElement('w:' + side)
                        side_el.set(qn('w:w'), '85')
                        side_el.set(qn('w:type'), 'dxa')
                        padding.append(side_el)
                    cell._tc.get_or_add_tcPr().append(padding)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.08
                    p.paragraph_format.widow_control = True
                    if row_i == 0:
                        p.paragraph_format.keep_with_next = True
                    if table_n == 3 and col_i < 2:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_inline(p, value)
                    for run in p.runs:
                        run.font.size = Pt(9.5)
                        if row_i == 0:
                            run.bold = True
                    if row_i == 0:
                        shade = OxmlElement('w:shd')
                        shade.set(qn('w:fill'), 'E7E6E6')
                        cell._tc.get_or_add_tcPr().append(shade)
                visible_rows.append([plain(v) for v in row])
            expected.append(('table', visible_rows))
            continue
        heading = re.match(r'^(#{1,4})\s+(.*)', line)
        if heading:
            level, text = len(heading[1]), heading[2]
            rendered = re.sub(r'[^\w\s]', ' ', plain(text))
            rendered = re.sub(r'\s+', ' ', rendered).strip()
            style = 'Title' if level == 1 else 'Heading ' + str(min(level - 1, 3))
            p = doc.add_paragraph(rendered, style)
            heading_map.append({'source': text, 'display': rendered})
            expected.append(('paragraph', rendered))
            continue
        if line.startswith('- '):
            line = line[2:]
            p = doc.add_paragraph(style='List Bullet')
        else:
            p = doc.add_paragraph()
        add_inline(p, line)
        if re.match(r'^\d+\. ', line):
            reference_n += 1
            for run in p.runs:
                run.font.size = Pt(10.5)
        expected.append(('paragraph', plain(line)))
    assert table_n == 3 and reference_n == 46
    doc.core_properties.title = plain(lines[0].lstrip('# '))
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.subject = 'Scientific reading draft with incomplete cohort analysis retained'
    QA.mkdir(parents=True)
    doc.save(OUTPUT)
    with ZipFile(OUTPUT) as archive:
        assert archive.testzip() is None
        root = ET.fromstring(archive.read('word/document.xml'))
        observed = []
        for el in root.find(W + 'body'):
            if el.tag == W + 'p':
                observed.append(('paragraph', visible_xml(el)))
            elif el.tag == W + 'tbl':
                observed.append(('table', [[visible_xml(tc) for tc in tr.findall(W + 'tc')] for tr in el.findall(W + 'tr')]))
        assert observed == expected, 'Visible text/table content changed'
        math_count = len(root.findall('.//' + M + 'oMath'))
        assert math_count == 2
        rels = ET.fromstring(archive.read('word/_rels/document.xml.rels'))
        hyperlinks = [r for r in rels if r.get('Type', '').endswith('/hyperlink')]
        assert {r.get('Target') for r in hyperlinks} == {r['target'] for r in links}
    assert SOURCE.read_bytes() == source
    result = {'status': 'CONTENT_PARITY_PASS_LAYOUT_PENDING', 'source_sha256': SOURCE_SHA,
              'output_bytes': OUTPUT.stat().st_size, 'output_sha256': hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
              'tables': table_n, 'reference_entries': reference_n, 'native_equations': math_count,
              'body_blocks_compared': len(expected), 'links': links, 'heading_display_mapping': heading_map,
              'science_text_unchanged': True, 'source_unchanged': True, 'figures_linked_not_embedded': True,
              'author_placeholders_retained': True, 'real_models_run': 0}
    with (QA / 'DOCX_CONTENT_CHECK.json').open('x', encoding='utf-8') as stream:
        json.dump(result, stream, ensure_ascii=False, indent=2)
    print(json.dumps({k: v for k, v in result.items() if k not in ('links', 'heading_display_mapping')}, sort_keys=True))


if __name__ == '__main__':
    build()
