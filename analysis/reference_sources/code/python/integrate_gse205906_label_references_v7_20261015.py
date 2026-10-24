"""Bounded v7 reporting successor, retaining v6 and all scientific outputs.

Only local prose and two new links change. Large computation remains on cloud;
compact source tables are read from R. No statistical producer is imported.
"""
import argparse
from copy import deepcopy
import csv
import hashlib
import importlib.util
from io import BytesIO
import json
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
MD = ROOT / "manuscript/Route_B2_scientific_reading_draft_en_20261014_integrated_v6.md"
WORD = MD.with_suffix(".docx")
FORMATTER = ROOT / "code/python/build_route_b2_reading_docx_20261010.py"
REC = ROOT / "results/diagnostics/GSE205906_gene_method_label_reference_20261015/INTEGRATION_RECOMMENDATION.md"
KREPORT = ROOT / "results/diagnostics/GSE205906_canonical_label_reference_20261015/SCIENCE_REPORT.md"
GREPORT = ROOT / "results/diagnostics/GSE205906_gene_method_label_reference_20261015/report_r02/SCIENCE_REPORT.md"
KR = Path("R:/Shaoyao_storage/cache/GSE205906_canonical_label_reference_20261015_r01")
GR = Path("R:/Shaoyao_storage/cache/GSE205906_gene_method_label_reference_20261015_r01")
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
      "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}


def require(ok, label):
    if not ok:
        raise ValueError(label)


def pin(path):
    path = Path(path)
    raw = path.read_bytes()
    return {"path": path.resolve().as_posix(), "bytes": len(raw),
            "sha256": hashlib.sha256(raw).hexdigest()}


def rows(path):
    with path.open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream, delimiter="\t"))


def package(path):
    with ZipFile(path) as archive:
        require(archive.testzip() is None, "DOCX CRC")
        return {name: archive.read(name) for name in archive.namelist()}


def canon(element):
    return etree.tostring(element, method="c14n")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-md", type=Path, required=True)
    parser.add_argument("--output-docx", type=Path, required=True)
    parser.add_argument("--qa", type=Path, required=True)
    args = parser.parse_args()
    require(not any(p.exists() for p in (args.output_md, args.output_docx, args.qa)), "Fresh outputs required")
    require(args.qa.resolve().as_posix().startswith("R:/Shaoyao_storage/"), "QA stays on R")
    fixed = {
        MD: "00296d85f9f8297a64146fe088b63016d4d2ab7a0f6f5f4792d49f3261b64499",
        WORD: "ad0c522122902497485d0d7edb03c5b8c4e7d460aec307eca2cf6a12b9a72f35",
        FORMATTER: "cdd11024f919a758c82fde8774d7f5cb27ad5505b44c6d57a738488eb2f3e7f8",
        REC: "204a9a128407269e7399cd8453b0fa6ca8c0f74263a7bc8f6298f6637c2d678c",
        KREPORT: "b6ae7547bdd362b8a526dbe28eeae7a54ce97aa784a14fdcd52e759f3011b7e7",
        GREPORT: "9df5505c8071a15765d3ea5bbc6b6e344691270f58c80058b894335ad877ae2e",
        KR / "ENUMERATION_COMPLETE.json": "3ae21310e0f06d23aa8b1163c8acc423551a62385a17fcb17bbfd108a8ad1015",
        GR / "ENUMERATION_COMPLETE.json": "4b3b7a16b73085d353dfd6dd7fe24928d719d23a608a7d091d0dc9d642fbce7f",
    }
    for report_dir, qa_sha in ((KREPORT.parent, "9f53d57769e72567e4e2a487cfa51db0147eb5d68b32171cdea0005a10fe67e9"),
                               (GREPORT.parent, "95835363337eb74ff20d9771070203be0ea6a2ced40f5033384785b0a4bddea4")):
        fixed[report_dir / "FINAL_REPORT_QA.md"] = qa_sha
    for directory in (KR, GR):
        completion = json.loads((directory / "ENUMERATION_COMPLETE.json").read_bytes())
        require(completion["allocations"] == 1680 and completion["pairwise_allocations"] == [20, 20]
                and completion["scoped_numerical_acceptance"] is True, "Complete reference required")
        for role in ("features", "family_summary"):
            identity = completion["outputs"][role]
            fixed[directory / Path(identity["path"]).name] = identity["sha256"]
    bindings = [pin(p) for p in fixed]
    require(all(b["sha256"] == fixed[Path(b["path"])] for b in bindings), "Pinned source mismatch")

    audit = []
    for directory, expected_rows, expected_families in ((KR, 31872, 10), (GR, 56868, 4)):
        features = rows(directory / "FEATURE_EMPIRICAL_REFERENCE.tsv")
        family_keys = {(r["collection"], r["method"], r["contrast"]) for r in features}
        require(len(features) == expected_rows and len(family_keys) == expected_families, "Full family inventory")
        require(all(int(r["global_N"]) == 1680 and int(r["pairwise_N"]) == 20 for r in features), "Tail denominators")
        require(all(float(r[col]) >= .05 for r in features for col in ("empirical_BH_global", "empirical_BH_pairwise")), "Empirical BH zero assertion")
        require(abs(min(float(r["exact_P_global"]) for r in features) - 1/840) < 1e-12
                and abs(min(float(r["exact_P_pairwise"]) for r in features) - .1) < 1e-12, "Attainable minima")
        for key in sorted(family_keys):
            family = [r for r in features if (r["collection"], r["method"], r["contrast"]) == key]
            audit.append({"claim": "/".join(key), "source": str(directory / "FEATURE_EMPIRICAL_REFERENCE.tsv"),
                          "family_n": len(family), "empirical_BH_global_lt_005": 0,
                          "empirical_BH_companion_lt_005": 0, "global_N": 1680, "companion_N": 20,
                          "verification": "Current compact accepted-output readback; not model refit"})
    ksummary = rows(KR / "FAMILY_REFERENCE_SUMMARY.tsv")
    gsummary = rows(GR / "FAMILY_REFERENCE_SUMMARY.tsv")
    displayed = []
    for collection, median, any_bh in (("Reactome", "16.7529", "88.5714"),
                                       ("SynGO_sensitivity", "24.3590", "83.5714")):
        selected = [r for r in ksummary if r["space"] == "global" and r["collection"] == collection]
        require(len(selected) == 2 and all(r["conditional_tail_warning"] == "TRUE" for r in selected), "Persistent CAMERA warning")
        require(all(f"{100*float(r['median_tail_05']):.4f}" == median and
                    f"{100*float(r['fraction_allocations_any_BH']):.4f}" == any_bh for r in selected), "CAMERA percentages")
        require(all(r["conditional_tail_warning"] == "TRUE" for r in ksummary if r["collection"] == collection), "Companion CAMERA warning")
        displayed.append({"collection": collection, "median_nominal_p005_percent": median, "any_native_BH_percent": any_bh})
    require(all(r["conditional_tail_warning"] == "FALSE" for r in gsummary), "Narrow gene-method warnings")
    for method, percentage in (("limma_voom_robust_eBayes", "11.9048"), ("DESeq2_Wald_sensitivity", "83.3333")):
        selected = [r for r in gsummary if r["space"] == "global" and r["method"] == method]
        require(len(selected) == 2 and all(f"{100*float(r['fraction_allocations_any_BH']):.4f}" == percentage for r in selected), "Gene-method percentages")
        displayed.append({"method": method, "any_native_BH_percent": percentage})

    proposed = REC.read_text(encoding="utf-8")
    paragraphs = {}
    for number in range(1, 7):
        match = re.search(rf"^### {number}\. [^\n]+\n\n(.+?)(?=\n\n)", proposed, flags=re.M | re.S)
        require(match is not None, "One recommendation paragraph per slot")
        paragraphs[number] = match[1]
    # Execution status only: no unfinished unsigned result is integrated.
    old = "Within GSE205906, genuine mroast/fgsea method-specific label references remain unrun at this reporting snapshot;"
    new = ("At the 15 October 2026 reporting snapshot, GSE205906 mroast/fgsea method-specific label references were running after observed compatibility checks, but their complete results and independent full-space checks were not yet available;")
    require(paragraphs[5].count(old) == 1, "Dated unsigned status slot")
    paragraphs[5] = paragraphs[5].replace(old, new, 1)
    # Remove compressed procedural language; the source meaning is unchanged.
    paragraphs[4] = paragraphs[4].replace("persistent CAMERA warnings further restrict evidence upgrading", "persistent CAMERA warnings further limit interpretation")
    paragraphs[6] = paragraphs[6].replace("scope of AI-assisted independent readbacks", "scope of separately implemented AI-assisted readbacks")

    text = MD.read_text(encoding="utf-8")
    source_paragraphs = text.split("\n\n")
    def slot(prefix):
        selected = [p for p in source_paragraphs if p.startswith(prefix)]
        require(len(selected) == 1, "Unique Markdown slot: " + prefix)
        return selected[0]
    old_design = slot("We conducted a secondary analysis")
    design = old_design.replace("Those results retain the 7 October reporting snapshot", "The original September/October effect reports retain the 7 October reporting snapshot; the conditional label-reference extensions are separately dated", 1)
    old_gate = slot("Only S1 had non-empty, fully evaluated families")
    gate = old_gate.replace("They did not complete unconditional source-wide RNA-seq label-null calibration, evaluate an admitted candidate or reopen upstream eligibility", "The score baselines did not themselves supply label-reference calibration; the subsequently completed GSE205906 conditional references cover specified methods only and do not complete the broader source-wide requirements, evaluate an admitted candidate or reopen upstream eligibility", 1)
    replacements = [(old_design, design, "Methods 2.1"),
                    (old_gate, gate, "Results 3.7"),
                    (slot("The new baseline comparisons distinguish"), paragraphs[4], "Discussion 4.4"),
                    (slot("The completed sensitivities and the three new score-baseline nodes"), paragraphs[5], "Discussion 4.5")]
    insertions = [
        ("### 2.6 Empty families, validation and reproducibility", [("#### 2.5.2 Retrospective conditional label references", True), (paragraphs[1], False)], "Methods 2.5.2"),
        (slot("[Figure 6: same-data method sensitivity"), [("#### 3.6.2 Complete conditional reference results", True), (paragraphs[2], False), (paragraphs[3], False)], "Results 3.6.2"),
        (slot("The [12 September GSE205906 working addendum]"), [(paragraphs[6], False)], "Supplementary materials"),
    ]
    for previous, revised, _ in replacements:
        require(previous != revised and text.count(previous) == 1, "Bounded replacement")
        text = text.replace(previous, revised, 1)
    for anchor, blocks, _ in insertions:
        require(text.count(anchor) == 1, "Bounded insertion")
        text = text.replace(anchor, "\n\n".join([t for t, _ in blocks]) + "\n\n" + anchor, 1)

    spec = importlib.util.spec_from_file_location("retained_inline_formatter", FORMATTER)
    fmt = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(fmt)
    fmt.SOURCE, fmt.links = args.output_md, []
    doc = Document(WORD)
    original = package(WORD)
    old_body = list(doc.element.body)
    old_xml = {id(el): canon(el) for el in old_body}
    removed, created, changes = set(), [], []
    body_template = next(p for p in doc.paragraphs if p.text.startswith("We conducted a secondary analysis"))
    heading_template = next(p for p in doc.paragraphs if p.text.startswith("2 5 1 Separately versioned"))
    def find(md):
        value = re.sub(r"^#{1,6} ", "", fmt.plain(md))
        if md.startswith("#"):
            value = re.sub(r"\s+", " ", re.sub(r"[^A-Za-z0-9 ]", " ", value)).strip()
        selected = [p for p in doc.paragraphs if p.text == value]
        require(len(selected) == 1, "Unique Word paragraph: " + md[:80])
        return selected[0]
    def make(md, template, heading=False):
        paragraph = doc.add_paragraph()
        ppr = template._p.find("w:pPr", NS)
        if ppr is not None:
            paragraph._p.append(deepcopy(ppr))
        value = md.removeprefix("#### ") if heading else md
        if heading:
            value = re.sub(r"\s+", " ", re.sub(r"[^A-Za-z0-9 ]", " ", value)).strip()
        fmt.add_inline(paragraph, value)
        require(paragraph.text == fmt.plain(value), "Word / Markdown text alignment")
        created.append(paragraph._p)
        return paragraph
    for previous, revised, label in replacements:
        p = find(previous)
        new_p = make(revised, p)
        removed.add(id(p._p))
        p._p.getparent().replace(p._p, new_p._p)
        changes.append({"location": label, "operation": "replace", "old": previous, "new": revised})
    for anchor, blocks, label in insertions:
        p = find(anchor)
        for value, heading in blocks:
            new_p = make(value, heading_template if heading else body_template, heading)
            p._p.addprevious(new_p._p)
        changes.append({"location": label, "operation": "insert_before", "anchor": anchor, "new": [v for v, _ in blocks]})
    untouched = [el for el in old_body if id(el) not in removed]
    require([el for el in doc.element.body if el not in created] == untouched, "Untouched body order")
    require(all(canon(el) == old_xml[id(el)] for el in untouched), "Untouched body XML")
    memory = BytesIO()
    doc.save(memory)
    with ZipFile(memory) as archive:
        generated = {n: archive.read(n) for n in archive.namelist()}
    editable = {"word/document.xml", "word/_rels/document.xml.rels"}
    require(set(original) == set(generated), "Original package inventory")
    for expression in (".//w:tbl", ".//m:oMath", ".//wp:inline", ".//w:sectPr"):
        a = etree.fromstring(original["word/document.xml"]).xpath(expression, namespaces=NS)
        b = etree.fromstring(generated["word/document.xml"]).xpath(expression, namespaces=NS)
        require([canon(e) for e in a] == [canon(e) for e in b], "Table equation image section preservation")
    old_rels = etree.fromstring(original["word/_rels/document.xml.rels"])
    new_rels = {e.get("Id"): canon(e) for e in etree.fromstring(generated["word/_rels/document.xml.rels"])}
    require(all(new_rels.get(e.get("Id")) == canon(e) for e in old_rels), "Old relationship preservation")
    require(all(pin(b["path"]) == b for b in bindings), "Source drift")
    args.qa.mkdir(parents=True)
    with args.output_md.open("x", encoding="utf-8", newline="\n") as stream:
        stream.write(text)
    with ZipFile(WORD) as archive, ZipFile(args.output_docx, "x") as output:
        for entry in archive.infolist():
            output.writestr(entry, generated[entry.filename] if entry.filename in editable else original[entry.filename])
    final = package(args.output_docx)
    require(all(final[n] == original[n] for n in set(original) - editable), "Preserve-only package parts")
    with (args.qa / "NUMERIC_CLAIM_READBACK.json").open("x", encoding="utf-8") as stream:
        json.dump({"families": audit, "displayed_percentages": displayed,
                   "scope": "Compact accepted output readback; not new model fits or full-space revalidation"}, stream, indent=2)
    record = {"status": "BOUNDED_CONTENT_AND_PACKAGE_CHECK_PASS_RENDER_PENDING",
              "created_utc": __import__("datetime").datetime.now(__import__("datetime").timezone.utc).isoformat(),
              "inputs": bindings, "builder": pin(__file__), "output_md": pin(args.output_md),
              "output_docx": pin(args.output_docx), "changes": changes,
              "numeric_readback": pin(args.qa / "NUMERIC_CLAIM_READBACK.json"),
              "untouched_body_elements_exact": len(untouched),
              "preserve_only_parts_exact": sorted(set(original) - editable),
              "all_prior_relationships_exact": True, "new_links": fmt.links,
              "tables_figures_equations_sections_exact": True,
              "old_files_modified_moved_deleted": False, "scientific_models_run": False,
              "unsigned_full_space_results_integrated": False,
              "whole_study_status": "NEEDS REVIEW", "submission_ready": False}
    with (args.qa / "INTEGRATION_CONTENT_CHECK.json").open("x", encoding="utf-8") as stream:
        json.dump(record, stream, indent=2, ensure_ascii=False)
        stream.write("\n")
    print(json.dumps({"status": record["status"], "output": record["output_docx"],
                      "local_slots": len(changes), "numeric_families": len(audit)}))


if __name__ == "__main__":
    main()
